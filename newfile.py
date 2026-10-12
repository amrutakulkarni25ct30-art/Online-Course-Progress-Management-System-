
import os
from openpyxl import Workbook, load_workbook

FILE = "/storage/emulated/0/course_progress.xlsx"

def setup():
    if not os.path.exists(FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Courses"
        ws.append([
            "Student ID", "Student Name",
            "Course Name", "Progress", "Status"
        ])
        wb.save(FILE)

def add_course():
    sid = input("Student ID: ")
    name = input("Student Name: ")
    course = input("Course Name: ")

    try:
        progress = int(input("Progress (0-100): "))
        if not 0 <= progress <= 100:
            print("Progress must be 0 to 100.")
            return
    except ValueError:
        print("Enter a valid number.")
        return

    status = "Completed" if progress == 100 else "In Progress"

    wb = load_workbook(FILE)
    ws = wb["Courses"]
    ws.append([sid, name, course, progress, status])
    wb.save(FILE)
    print("Course added successfully!")

def view_courses():
    wb = load_workbook(FILE)
    ws = wb["Courses"]

    print("\n--- ALL COURSE RECORDS ---")
    for row in ws.iter_rows(min_row=2, values_only=True):
        print(row)

def update_progress():
    sid = input("Student ID: ")
    course = input("Course Name: ")

    try:
        progress = int(input("New Progress (0-100): "))
        if not 0 <= progress <= 100:
            print("Progress must be 0 to 100.")
            return
    except ValueError:
        print("Enter a valid number.")
        return

    wb = load_workbook(FILE)
    ws = wb["Courses"]

    for row in range(2, ws.max_row + 1):
        if (str(ws.cell(row, 1).value) == sid and
            str(ws.cell(row, 3).value).lower() == course.lower()):

            ws.cell(row, 4).value = progress
            ws.cell(row, 5).value = (
                "Completed" if progress == 100
                else "In Progress"
            )
            wb.save(FILE)
            print("Progress updated!")
            return

    print("Course record not found.")

def main():
    setup()

    while True:
        print("\n=== ONLINE COURSE PROGRESS SYSTEM ===")
        print("1. Add Course")
        print("2. View Courses")
        print("3. Update Progress")
        print("4. Exit")

        choice = input("Choose option: ")

        if choice == "1":
            add_course()
        elif choice == "2":
            view_courses()
        elif choice == "3":
            update_progress()
        elif choice == "4":
            print("Thank you!")
            break
        else:
            print("Invalid choice.")

if __name__ == "__main__":
    setup()
    main()
